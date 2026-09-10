"""
extract.py — Raízes Dashboard Data Extractor
Reads raizes_data_fresh.xlsx and produces the DATA JSON block
embedded inside raizes_dashboard.html.
"""

import openpyxl
import datetime
import json
import re
from collections import defaultdict

XLSX_PATH = "/home/claude/raizes_data_fresh.xlsx"
SALDO_INICIAL_2024 = 83112.61

MESES_PT = ['jan','fev','mar','abr','mai','jun','jul','ago','set','out','nov','dez']
MESES_MAP = {m.upper(): i+1 for i,m in enumerate(MESES_PT)}
MESES_MAP.update({'JANEIRO':1,'FEVEREIRO':2,'MARÇO':3,'ABRIL':4,'MAIO':5,'JUNHO':6,
                  'JULHO':7,'AGOSTO':8,'SETEMBRO':9,'OUTUBRO':10,'NOVEMBRO':11,'DEZEMBRO':12})

EXCLUDED_CATS = {'Ajuste Contábil','Dona do Meu Fluxo','Dona do Meu Fluxo '}
CAT_ALIASES = {'Despesas com equipe':'Despesas com Equipe','Despesas com equipe ':'Despesas com Equipe',
               'Despesas com Equipe ':'Despesas com Equipe','Projetos ':'Projetos',
               'Administrativo ':'Administrativo','Financeiro ':'Financeiro',
               'Dona do Meu Fluxo':'Ajuste Contábil'}
PROJ_ALIASES = {'Almenara ':'Almenara','Maria da Fé ':'Maria da Fé','re.green 2':'Re.green 2',
                're.green2':'Re.green 2','Re.green2':'Re.green 2','sta barbara':'Santa Barbara',
                'Sta Barbara':'Santa Barbara','trs arbo':'TRs Arbo','Trs Arbo':'TRs Arbo',
                'unesco':'UNESCO','Unesco':'UNESCO','vli':'VLI'}

def norm_cat(v):
    if not v: return None
    s = str(v).strip()
    return CAT_ALIASES.get(s, s)

def norm_proj(v):
    if not v: return None
    s = str(v).strip()
    return PROJ_ALIASES.get(s, s)

def norm_str(v):
    if not v: return None
    return str(v).strip() or None

def parse_date_cell(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, str):
        m = re.match(r'(\d{2})/(\d{2})/(\d{4})', v.strip())
        if m: return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None

def parse_mmyy(v):
    if not v: return None
    s = str(v).strip()
    m = re.match(r'(\d{1,2})/(\d{2,4})', s)
    if not m: return None
    mo, yr = int(m.group(1)), int(m.group(2))
    if yr < 100: yr += 2000
    return datetime.date(yr, mo, 1)

def label(ym):
    year, month = ym.split('-')
    return f"{MESES_PT[int(month)-1]}/{year[2:]}"

def floatv(v, default=0.0):
    if v is None: return default
    try: return float(v)
    except: return default

wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

# ── 1. LANÇAMENTOS ──
ws_lanc = wb['Lançamentos']
raw_rows = list(ws_lanc.iter_rows(min_row=5, values_only=True))

tx = []
tx_futuro = []

for r in raw_rows:
    d = parse_date_cell(r[0])
    cat = norm_cat(r[5])
    if not cat: continue
    credito = floatv(r[8])
    debito  = floatv(r[9])
    saldo   = floatv(r[10]) if r[10] is not None else None
    pagador = norm_str(r[3])
    detalhamento = norm_proj(r[6]) if cat == 'Projetos' else norm_str(r[6])
    det_proj = norm_str(r[7])
    local       = norm_str(r[13])
    minoria     = norm_str(r[14])
    tipo_desloc = norm_str(r[15])
    km          = floatv(r[16])

    if d:
        tx.append({'data':d,'cat':cat,'pagador':pagador,'detalhamento':detalhamento,
                   'det_proj':det_proj,'credito':credito,'debito':debito,'saldo':saldo,
                   'local':local,'minoria':minoria,'tipo_desloc':tipo_desloc,'km':km})
    else:
        mes_str = norm_str(r[2])
        if not mes_str: continue
        mes_num = MESES_MAP.get(mes_str.upper())
        if not mes_num: continue
        if not pagador and not credito and not debito: continue
        tx_futuro.append({'mes_num':mes_num,'cat':cat,'pagador':pagador,'detalhamento':detalhamento,
                          'det_proj':det_proj,'credito':credito,'debito':debito})

tx.sort(key=lambda x: x['data'])
tx_incl = [t for t in tx if t['cat'] not in EXCLUDED_CATS]
tx_futuro_incl = [t for t in tx_futuro if t['cat'] not in EXCLUDED_CATS]

last_real_date  = tx[-1]['data']
last_real_month = last_real_date.month

def futuro_ym(mes_num):
    year = last_real_date.year if mes_num >= last_real_month else last_real_date.year + 1
    return f"{year}-{mes_num:02d}"

# ── 2. MONTHLY ──
monthly_dict = defaultdict(lambda: {'creditos':0.0,'debitos':0.0,'saldo_fim':None,'previsto':False})
for t in tx_incl:
    ym = f"{t['data'].year}-{t['data'].month:02d}"
    monthly_dict[ym]['creditos'] += t['credito']
    monthly_dict[ym]['debitos']  += t['debito']
for t in tx:
    if t['saldo'] is not None:
        ym = f"{t['data'].year}-{t['data'].month:02d}"
        monthly_dict[ym]['saldo_fim'] = t['saldo']

monthly = []
for ym in sorted(monthly_dict.keys()):
    d = monthly_dict[ym]
    monthly.append({'ym':ym,'label':label(ym),'creditos':round(d['creditos'],2),
                    'debitos':round(d['debitos'],2),'resultado':round(d['creditos']+d['debitos'],2),
                    'saldo_fim':round(d['saldo_fim'],2) if d['saldo_fim'] is not None else None,'previsto':False})

prev_saldo = SALDO_INICIAL_2024
for m in monthly:
    if m['saldo_fim'] is None: m['saldo_fim'] = round(prev_saldo + m['resultado'],2)
    prev_saldo = m['saldo_fim']

meses_disponiveis = [m['ym'] for m in monthly]

# ── 3. MONTHLY FUTURO ──
futuro_dict = defaultdict(lambda: {'creditos':0.0,'debitos':0.0})
for t in tx_futuro_incl:
    ym = futuro_ym(t['mes_num'])
    futuro_dict[ym]['creditos'] += t['credito']
    futuro_dict[ym]['debitos']  += t['debito']

meses_futuros = sorted(futuro_dict.keys())
monthly_futuro = list(monthly)

for ym in meses_futuros:
    existing = next((m for m in monthly_futuro if m['ym']==ym), None)
    fd = futuro_dict[ym]
    if existing:
        idx2 = monthly_futuro.index(existing)
        new_c = round(existing['creditos']+fd['creditos'],2)
        new_d = round(existing['debitos']+fd['debitos'],2)
        monthly_futuro[idx2] = {**existing,'creditos':new_c,'debitos':new_d,'resultado':round(new_c+new_d,2),'previsto':True}
    else:
        prev_s = monthly_futuro[-1]['saldo_fim'] if monthly_futuro else SALDO_INICIAL_2024
        c,d2 = round(fd['creditos'],2), round(fd['debitos'],2)
        r = round(c+d2,2)
        monthly_futuro.append({'ym':ym,'label':label(ym),'creditos':c,'debitos':d2,'resultado':r,
                               'saldo_fim':round(prev_s+r,2),'previsto':True})

prev_saldo = None
for m in monthly_futuro:
    if not m['previsto']: prev_saldo = m['saldo_fim']
    else:
        if prev_saldo is not None: m['saldo_fim'] = round(prev_saldo+m['resultado'],2)
        prev_saldo = m['saldo_fim']

# ── 4. LANÇAMENTOS POR MÊS ──
lanc_por_mes = defaultdict(list)
for t in tx:
    ym = f"{t['data'].year}-{t['data'].month:02d}"
    lanc_por_mes[ym].append({'data':t['data'].isoformat(),'pagador':t['pagador'],
        'categoria':t['cat'],'detalhamento':t['detalhamento'],'det_proj':t['det_proj'],
        'credito':t['credito'],'debito':t['debito'],'saldo':t['saldo'],'previsto':False})
for t in tx_futuro:
    ym = futuro_ym(t['mes_num'])
    lanc_por_mes[ym].append({'data':None,'pagador':t['pagador'],'categoria':t['cat'],
        'detalhamento':t['detalhamento'],'det_proj':t['det_proj'],
        'credito':t['credito'],'debito':t['debito'],'saldo':None,'previsto':True})

# ── 5. CATEGORIAS POR MÊS ──
def build_cat_por_mes(transactions, month_list, is_futuro=False):
    cat_det = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: {'credito':0.0,'debito':0.0})))
    for t in transactions:
        ym = futuro_ym(t['mes_num']) if is_futuro else f"{t['data'].year}-{t['data'].month:02d}"
        det = t['detalhamento'] or 'Outros'
        cat_det[ym][t['cat']][det]['credito'] += t['credito']
        cat_det[ym][t['cat']][det]['debito']  += t['debito']
    result = {}
    for ym in month_list:
        result[ym] = {}
        for cat, dets in cat_det[ym].items():
            total_c = sum(v['credito'] for v in dets.values())
            total_d = sum(v['debito']  for v in dets.values())
            detalhes = sorted([{'nome':det,'credito':round(v['credito'],2),'debito':round(v['debito'],2)}
                                for det,v in dets.items()], key=lambda x: abs(x['debito']), reverse=True)
            result[ym][cat] = {'total_credito':round(total_c,2),'total_debito':round(total_d,2),'detalhes':detalhes}
    return result

categorias_por_mes        = build_cat_por_mes(tx_incl, meses_disponiveis)
categorias_por_mes_futuro = build_cat_por_mes(tx_futuro_incl, meses_futuros, is_futuro=True)

# ── 6. IMPOSTOS ──
ws_imp = wb['Impostos']
tax_rate = {}
for r in list(ws_imp.iter_rows(min_row=2, values_only=True)):
    if r[0] and r[1] and r[2]:
        tax_rate[f"{int(r[0])}-{int(r[1]):02d}"] = floatv(r[2])

def get_rate(ym):
    if ym in tax_rate: return tax_rate[ym]
    rates_before = [(k,v) for k,v in tax_rate.items() if k<=ym]
    return max(rates_before, key=lambda x: x[0])[1] if rates_before else 0.0

# ── 7. PROJETOS ──
def build_project_financials(transactions):
    proj_data = defaultdict(lambda: {'entradas':0.0,'gastos':0.0,'detalhes':defaultdict(lambda:{'debito':0.0,'credito':0.0})})
    org_rev_by_month  = defaultdict(float)
    proj_rev_by_month = defaultdict(lambda: defaultdict(float))
    for t in transactions:
        if t['cat'] != 'Projetos': continue
        proj = t['detalhamento'] or 'Sem Projeto'
        det  = t['det_proj'] or 'Outros'
        ym   = futuro_ym(t['mes_num']) if t.get('mes_num') else f"{t['data'].year}-{t['data'].month:02d}"
        c,d  = t['credito'], t['debito']
        proj_data[proj]['entradas'] += c
        proj_data[proj]['gastos']   += d
        proj_data[proj]['detalhes'][det]['credito'] += c
        proj_data[proj]['detalhes'][det]['debito']  += d
        if c > 0:
            org_rev_by_month[ym]        += c
            proj_rev_by_month[ym][proj] += c
    proj_tax = defaultdict(float)
    for ym, org_rev in org_rev_by_month.items():
        if org_rev == 0: continue
        total_tax = org_rev * get_rate(ym)
        for proj, proj_rev in proj_rev_by_month[ym].items():
            proj_tax[proj] += total_tax * (proj_rev/org_rev)
    result = []
    for proj, d in sorted(proj_data.items(), key=lambda x: x[1]['entradas'], reverse=True):
        entradas=round(d['entradas'],2); gastos=round(d['gastos'],2)
        impostos_est=round(-proj_tax[proj],2); saldo=round(entradas+gastos+impostos_est,2)
        margem=round((saldo/entradas*100) if entradas>0 else 0,1)
        detalhes=sorted([{'nome':k,'debito':round(v['debito'],2),'credito':round(v['credito'],2)}
                          for k,v in d['detalhes'].items()], key=lambda x: abs(x['debito']), reverse=True)
        result.append({'nome':proj,'entradas':entradas,'gastos':gastos,'impostos_est':impostos_est,
                       'saldo':saldo,'margem':margem,'detalhes':detalhes})
    return result

projetos           = build_project_financials(tx_incl)
projetos_com_futuro= build_project_financials(tx_incl + tx_futuro_incl)

# ── 8. PROPOSTA DOS PROJETOS ──
ws_prop = wb['Proposta dos Projetos']
prop_rows = list(ws_prop.iter_rows(min_row=4, values_only=True))
proj_dates  = {}
proj_budget = defaultdict(float)
for r in prop_rows:
    nome_date  = norm_proj(r[13]) if len(r)>13 else None
    inicio     = parse_mmyy(r[14]) if len(r)>14 else None
    fim        = parse_mmyy(r[15]) if len(r)>15 else None
    val_total  = floatv(r[16]) if len(r)>16 else 0
    if nome_date and nome_date not in ('Nome Projeto','Nome do Projeto') and inicio and fim:
        proj_dates[nome_date] = {'inicio':inicio,'fim':fim,'valor_total':val_total}
    nome_budget = norm_proj(r[0]) if r[0] else None
    val = floatv(r[8]) if len(r)>8 else 0
    if nome_budget and nome_budget not in ('Nome Projeto','Nome do Projeto','') and val:
        proj_budget[nome_budget] += val

# ── 9. RITMO DOS PROJETOS ──
today = datetime.date.today()
gastos_by_proj   = defaultdict(float)
entradas_by_proj = defaultdict(float)
for t in tx_incl:
    if t['cat']=='Projetos' and t['detalhamento']:
        gastos_by_proj[t['detalhamento']]   += abs(t['debito'])
        entradas_by_proj[t['detalhamento']] += t['credito']

def ritmo_status(tp, gp):
    if tp>=100: return 'encerrado'
    if gp-tp>15: return 'critico'
    if gp-tp>5:  return 'atencao'
    return 'no_prazo'

ritmo_projetos = []
for proj in sorted(set(list(proj_dates.keys()) + list(proj_budget.keys()))):
    pd = proj_dates.get(proj)
    if not pd: continue
    inicio=pd['inicio']; fim=pd['fim']
    fim_end = datetime.date(fim.year,fim.month,
        (datetime.date(fim.year+(fim.month//12),(fim.month%12)+1,1)-datetime.timedelta(days=1)).day if fim.month<12 else 31)
    duracao  = (fim_end-inicio).days or 1
    decorrido= max(0,min((today-inicio).days,duracao))
    tempo_pct= round(decorrido/duracao*100,1)
    planejado= proj_budget.get(proj, pd.get('valor_total',0))
    gasto    = round(gastos_by_proj.get(proj,0),2)
    gasto_pct= round(gasto/planejado*100,1) if planejado>0 else 0.0
    valor_total = pd.get('valor_total') or 0
    entradas = round(entradas_by_proj.get(proj,0),2)
    receb_pct = round(entradas/valor_total*100,1) if valor_total>0 else 0.0
    ritmo_projetos.append({'nome':proj,'status':ritmo_status(tempo_pct,gasto_pct),
        'tempo_pct':tempo_pct,'gasto_pct':gasto_pct,'planejado':round(planejado,2),'gasto':gasto,
        'entradas':entradas,'valor_total':round(valor_total,2),'receb_pct':receb_pct,
        'inicio':inicio.isoformat(),'fim':fim_end.isoformat()})

# ── 10. PLANEJADO × REALIZADO ──
plan_by_det    = defaultdict(float)
real_by_det    = defaultdict(lambda:{'credito':0.0,'debito':0.0})
plan_by_proj_det = defaultdict(lambda: defaultdict(float))
real_by_proj_det = defaultdict(lambda: defaultdict(lambda:{'credito':0.0,'debito':0.0}))

for r in prop_rows:
    proj = norm_proj(r[0]) if r[0] else None
    det  = norm_str(r[2]) if len(r)>2 else None
    val  = floatv(r[8]) if len(r)>8 else 0
    if det and val:
        plan_by_det[det] += val
        if proj and proj not in ('Nome Projeto','Nome do Projeto',''):
            plan_by_proj_det[proj][det] += val

for t in tx_incl:
    if t['cat']=='Projetos':
        det = t['det_proj'] or 'Sem Categoria'
        real_by_det[det]['credito'] += t['credito']
        real_by_det[det]['debito']  += t['debito']
        if t['detalhamento']:
            real_by_proj_det[t['detalhamento']][det]['credito'] += t['credito']
            real_by_proj_det[t['detalhamento']][det]['debito']  += t['debito']

avg_tempo = (sum(p['tempo_pct']*p['planejado'] for p in ritmo_projetos) /
             sum(p['planejado'] for p in ritmo_projetos if p['planejado']>0)) if any(p['planejado'] for p in ritmo_projetos) else 50.0

all_dets = sorted(set(list(plan_by_det.keys()) + list(real_by_det.keys())))
planejado_realizado = sorted([{
    'nome':det,
    'planejado':round(plan_by_det.get(det,0),2),
    'planejado_momento':round(plan_by_det.get(det,0)*avg_tempo/100,2),
    'realizado':round(abs(real_by_det[det]['debito']),2),
    'saldo':round(plan_by_det.get(det,0)-abs(real_by_det[det]['debito']),2)
} for det in all_dets], key=lambda x: x['realizado'], reverse=True)

def build_planreal_proj(proj):
    rp = next((p for p in ritmo_projetos if p['nome']==proj),None)
    tp = rp['tempo_pct'] if rp else avg_tempo
    dets = sorted(set(list(plan_by_proj_det.get(proj,{}).keys())+list(real_by_proj_det.get(proj,{}).keys())))
    return sorted([{'nome':det,'planejado':round(plan_by_proj_det[proj].get(det,0),2),
        'planejado_momento':round(plan_by_proj_det[proj].get(det,0)*tp/100,2),
        'realizado':round(abs(real_by_proj_det[proj][det]['debito']),2),
        'saldo':round(plan_by_proj_det[proj].get(det,0)-abs(real_by_proj_det[proj][det]['debito']),2)
    } for det in dets], key=lambda x: x['planejado'], reverse=True)

all_proj_list = sorted(set(list(plan_by_proj_det.keys())+list(real_by_proj_det.keys())))
planejado_realizado_by_projeto = {'Todos':planejado_realizado}
for p in all_proj_list: planejado_realizado_by_projeto[p] = build_planreal_proj(p)

all_proj_names = set(proj_dates.keys())|set(proj_budget.keys())|set(gastos_by_proj.keys())
projetos_planreal = sorted([{
    'nome':proj,'planejado':round(proj_budget.get(proj,proj_dates.get(proj,{}).get('valor_total',0)),2),
    'planejado_momento':round(proj_budget.get(proj,0)*(next((p['tempo_pct'] for p in ritmo_projetos if p['nome']==proj),0))/100,2),
    'realizado':round(gastos_by_proj.get(proj,0),2),
    'saldo':round(proj_budget.get(proj,0)-gastos_by_proj.get(proj,0),2)
} for proj in sorted(all_proj_names)], key=lambda x: x['planejado'], reverse=True)

lista_projetos = sorted(all_proj_names)

# ── 11. INDICADORES BY PROJECT ──
ind_raw = defaultdict(lambda:{'local':defaultdict(lambda:{'valor':0.0}),
                               'minoria':defaultdict(lambda:{'valor':0.0}),
                               'desloc':defaultdict(lambda:{'debito':0.0,'km':0.0})})
for t in tx:
    if t['cat']!='Projetos' or not t['detalhamento']: continue
    proj=t['detalhamento']; spend=abs(t['debito'])
    lv = 'Sim' if str(t['local'] or '').lower()=='sim' else ('Não' if str(t['local'] or '').lower() in ('não','nao') else 'Não sei')
    mv = 'Sim' if str(t['minoria'] or '').lower()=='sim' else ('Não' if str(t['minoria'] or '').lower() in ('não','nao') else 'Não sei')
    ind_raw[proj]['local'][lv]['valor']   += spend
    ind_raw[proj]['minoria'][mv]['valor'] += spend
    if t['km'] and t['km']>0:
        tipo = str(t['tipo_desloc'] or '').strip() or 'Outros'
        ind_raw[proj]['desloc'][tipo]['debito'] += spend
        ind_raw[proj]['desloc'][tipo]['km']     += t['km']

indicadores_by_projeto = {}
for proj,d in ind_raw.items():
    tl=sum(v['valor'] for v in d['local'].values()) or 1
    tm=sum(v['valor'] for v in d['minoria'].values()) or 1
    td=sum(v['debito'] for v in d['desloc'].values()) or 1
    indicadores_by_projeto[proj]={
        'local':    sorted([{'nome':k,'valor':round(v['valor'],2),'pct':round(v['valor']/tl*100,1)} for k,v in d['local'].items()],    key=lambda x:-x['valor']),
        'minoria':  sorted([{'nome':k,'valor':round(v['valor'],2),'pct':round(v['valor']/tm*100,1)} for k,v in d['minoria'].items()],  key=lambda x:-x['valor']),
        'deslocamento': sorted([{'nome':k,'debito':round(v['debito'],2),'km':round(v['km'],1),'pct':round(v['debito']/td*100,1)} for k,v in d['desloc'].items()], key=lambda x:-x['debito']),
    }

# ── 12. GASTOS MENSAIS BY PROJECT ──
def build_gastos_mensais(transactions):
    result = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    for t in transactions:
        if t['cat']!='Projetos' or not t['detalhamento']: continue
        det = t['det_proj'] or 'Outros'
        ym  = futuro_ym(t['mes_num']) if t.get('mes_num') else f"{t['data'].year}-{t['data'].month:02d}"
        result[t['detalhamento']][det][ym] += abs(t['debito'])
    return {proj:{det:{ym:round(v,2) for ym,v in sorted(months.items())} for det,months in dets.items()} for proj,dets in result.items()}

gastos_mensais_by_projeto        = build_gastos_mensais(tx_incl)
gastos_mensais_futuro_by_projeto = {k:v for k,v in build_gastos_mensais(tx_futuro_incl).items() if v}

# ── 13. PREVISÃO 2026 ──
ws_prev = wb['Fluxo Previsto 2026']
CAT_MAP_PREV = {'RECEITAS DE VENDAS E SERVIÇOS':'Projetos','OUTRAS ENTRADAS':'Outras Entradas',
                'RECEITAS NÃO OPERACIONAIS':'Receitas Não Operacionais','DESPESAS ADMINISTRATIVAS':'Administrativo',
                'DESPESAS COM EQUIPE':'Despesas com Equipe','DESPESAS DE ESPAÇO FÍSICO':'Despesas de Espaço Físico',
                'DESPESAS FINANCEIRAS':'Taxas Bancárias','DESPESAS COMUNICAÇÃO':'Comunicação',
                'DESPESAS PROJETOS PRÓPRIOS':'Projetos','DESPESAS COM VENDAS':'Administrativo',
                'IMPOSTOS':'Impostos'}

# Map lancamento (categoria, detalhamento) → categoria_real from previsao
# Built by matching detalhamento in lancamentos to subcategory names in previsao
# Key insight: lancamento 'categoria' maps to CAT_MAP_PREV values, 'detalhamento' matches previsao subcategories
LANC_CAT_TO_PREV = {v:v for v in CAT_MAP_PREV.values()}  # identity for already-mapped

prev_by_catdet = defaultdict(lambda:[0.0]*12); prev_tipos={}; prev_cat_map={}
prev_subcat_to_cat = {}  # subcat name -> categoria_real
for r in list(ws_prev.iter_rows(min_row=2, values_only=True)):
    cat1=norm_str(r[1]); det=norm_str(r[4]); data_cell=r[5]; valor=floatv(r[6])
    if not cat1 or not data_cell or not valor: continue
    if not isinstance(data_cell, datetime.datetime): continue
    cat_mapped = CAT_MAP_PREV.get(cat1,cat1); key=f"{cat_mapped}||{det or cat_mapped}"
    prev_tipos[key] = 'entrada' if ('RECEITAS' in cat1 or 'ENTRADA' in cat1) else 'saida'
    prev_by_catdet[key][data_cell.month-1] += valor
    prev_cat_map[key] = cat_mapped
    if det: prev_subcat_to_cat[det] = cat_mapped

previsao_2026 = [{'categoria':key.split('||',1)[1],'categoria_real':prev_cat_map[key],'tipo':prev_tipos[key],
                   'meses':[round(v,2) for v in vals]} for key,vals in sorted(prev_by_catdet.items())]

# ── Aggregate real 2026 lancamentos by (previsao_cat_real, detalhamento, mes) ──
# Match lancamento detalhamento → previsao subcategory → cat_real + tipo
# Unmatched → 'Sem Categoria'
prev_sub_map = {}  # det_name -> (cat_real, tipo)
prev_sub_map.update({
    'ABETA':('Outras Entradas','entrada'), 'Agência de Comunicação':('Comunicação','saida'),
    'Ações Climáticas':('Administrativo','saida'), 'Bem Estar Equipe':('Despesas com Equipe','saida'),
    'Bem estar':('Despesas com Equipe','saida'),
    'Benefícios':('Despesas com Equipe','saida'), 'Coletivo MUDA':('Administrativo','saida'),
    'Combu':('Projetos','entrada'), 'Contabilidade':('Administrativo','saida'),
    'Designer':('Comunicação','saida'), 'Despesas Captação':('Administrativo','saida'),
    'Diretor':('Despesas com Equipe','saida'), 'Especialista':('Despesas com Equipe','saida'),
    'Eventos':('Administrativo','saida'), 'Eventos/Captação':('Administrativo','saida'),
    'FGTS':('Despesas com Equipe','saida'), 'HUB BH':('Despesas de Espaço Físico','saida'),
    'INSS':('Despesas com Equipe','saida'), 'IPTU':('Despesas de Espaço Físico','saida'),
    'IRRF':('Despesas com Equipe','saida'), 'Imersão':('Administrativo','saida'),
    'Manutenção':('Despesas de Espaço Físico','saida'), 'Material escritório':('Administrativo','saida'),
    'Materiais de Escritório':('Administrativo','saida'), 'NET':('Despesas de Espaço Físico','saida'),
    'Outros':('Administrativo','saida'), 'Outros Impostos':('Impostos','saida'),
    'Programas/Assinaturas':('Administrativo','saida'), 'Qualificação':('Despesas com Equipe','saida'),
    'Redes Sociais':('Comunicação','saida'), 'Rend.Aut':('Receitas Não Operacionais','entrada'),
    'Rend':('Receitas Não Operacionais','entrada'),
    'SIMPLES':('Impostos','saida'), 'Saúde e Segurança':('Administrativo','saida'),
    'Sistema B':('Administrativo','saida'), 'Site':('Comunicação','saida'),
    'Suporte TI':('Administrativo','saida'), 'São Paulo':('Despesas de Espaço Físico','saida'),
    'Tarifa':('Taxas Bancárias','saida'), 'Assistente':('Despesas com Equipe','saida'),
    # Projetos: entries are entradas, expenditures are saidas — split by sign
})
# Project names → entradas when credito, saidas when debito
PROJ_NAMES = set(norm_proj(r[0]) or '' for r in prop_rows if r[0]) - {''}

def get_cat_tipo_for_lanc(t):
    det = t['detalhamento'] or ''
    # Check direct subcategory match first
    if det in prev_sub_map:
        cat_real, tipo = prev_sub_map[det]
        # For entries that can be both (projects): use sign
        if cat_real == 'Projetos' and tipo == 'entrada' and t.get('debito',0) < 0:
            return (cat_real, 'saida')
        return (cat_real, tipo)
    # Project names: credit=entrada, debit=saida
    if det in PROJ_NAMES or t['cat'] == 'Projetos':
        return ('Projetos', 'entrada' if (t.get('credito',0) or 0) > 0 else 'saida')
    # Fallback by categoria
    CAT_FALLBACK = {
        'Administrativo':('Administrativo','saida'), 'Comunicação':('Comunicação','saida'),
        'Despesas com Equipe':('Despesas com Equipe','saida'),
        'Despesas com equipe':('Despesas com Equipe','saida'),
        'Despesas de Espaço Físico':('Despesas de Espaço Físico','saida'),
        'Impostos':('Impostos','saida'), 'Taxas Bancárias':('Taxas Bancárias','saida'),
        'Financeiro':('Financeiro','saida'),
    }
    return CAT_FALLBACK.get(t['cat'], ('Sem Categoria', 'saida' if (t.get('debito',0) or 0) < 0 else 'entrada'))

# {cat_real -> {tipo -> {det -> {'c':[12],'d':[12]}}}}
real_prev_2026   = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda:{'c':[0.0]*12,'d':[0.0]*12})))
futuro_prev_2026 = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda:{'c':[0.0]*12,'d':[0.0]*12})))

for t in tx:
    if t['data'].year != 2026: continue
    cat_real, tipo = get_cat_tipo_for_lanc(t)
    det = t['detalhamento'] or t['cat'] or 'Sem Categoria'
    mes = t['data'].month - 1
    real_prev_2026[cat_real][tipo][det]['c'][mes] += t['credito']
    real_prev_2026[cat_real][tipo][det]['d'][mes] += t['debito']

for t in tx_futuro:
    cat_real, tipo = get_cat_tipo_for_lanc(t)
    det = t['detalhamento'] or t['cat'] or 'Sem Categoria'
    mes_num = t['mes_num'] - 1
    futuro_prev_2026[cat_real][tipo][det]['c'][mes_num] += t['credito']
    futuro_prev_2026[cat_real][tipo][det]['d'][mes_num] += t['debito']

def ser_real_prev(d):
    # {cat_real -> {tipo -> {det -> {'c':[12],'d':[12]}}}}
    return {cat: {tipo: {det: {'c':[round(v,2) for v in cd['c']],'d':[round(v,2) for v in cd['d']]}
                         for det,cd in tipo_dets.items()}
                  for tipo,tipo_dets in tipos.items()}
            for cat,tipos in d.items()}

# ── 14. APLICAÇÃO ──
ws_aplic = wb['Aplicação']
aplicacao = []
for r in list(ws_aplic.iter_rows(min_row=2, values_only=True)):
    if not r[0] or not r[1]: continue
    try: ano=int(floatv(r[0])); mes=int(floatv(r[1]))
    except: continue
    aplicacao.append({'ano':ano,'mes':mes,'entrada':round(floatv(r[2]),2),'saida':round(floatv(r[3]),2),
                      'saldo_final':round(floatv(r[4]),2) if r[4] is not None else None,
                      'rendimento':round(floatv(r[5]),2)})

# ── ASSEMBLE ──
data_out = {
    'gerado_em': datetime.datetime.now().isoformat(timespec='seconds'),
    'periodo_dados': f"{tx[0]['data'].isoformat()} a {tx[-1]['data'].isoformat()}",
    'monthly':monthly,'monthly_futuro':monthly_futuro,
    'lancamentos_por_mes':dict(lanc_por_mes),
    'categorias_por_mes':categorias_por_mes,'categorias_por_mes_futuro':categorias_por_mes_futuro,
    'meses_disponiveis':meses_disponiveis,'meses_futuros':meses_futuros,
    'projetos':projetos,'projetos_com_futuro':projetos_com_futuro,
    'ritmo_projetos':ritmo_projetos,
    'planejado_realizado':planejado_realizado,
    'planejado_realizado_by_projeto':planejado_realizado_by_projeto,
    'planejado_realizado_by_projeto_futuro':planejado_realizado_by_projeto,
    'projetos_planreal':projetos_planreal,
    'indicadores_by_projeto':indicadores_by_projeto,
    'gastos_mensais_by_projeto':gastos_mensais_by_projeto,
    'gastos_mensais_futuro_by_projeto':gastos_mensais_futuro_by_projeto,
    'lista_projetos':lista_projetos,
    'previsao_2026':previsao_2026,
    'real_prev_2026':ser_real_prev(real_prev_2026),
    'futuro_prev_2026':ser_real_prev(futuro_prev_2026),
    'aplicacao':aplicacao,
}
print(json.dumps(data_out, ensure_ascii=False, default=str))
